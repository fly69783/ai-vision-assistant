from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

doc = Document()

# ===== 全局样式设置 =====
style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
font.size = Pt(11)
font.color.rgb = RGBColor(0x33, 0x33, 0x33)

# 段落间距
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.35

# ===== 封面标题 =====
doc.add_paragraph()  # 空行
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('AI 视觉辅助盲人环境理解系统')
run.bold = True
run.font.size = Pt(22)
run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
run.font.name = '微软雅黑'
run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('项目介绍与安排 —— 演讲逐字稿')
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
run.font.name = '微软雅黑'
run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

date_p = doc.add_paragraph()
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = date_p.add_run('2026年7月16日')
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

doc.add_paragraph()  # 空行

# ===== 辅助函数 =====
def add_section_title(text, level=1):
    """添加章节标题"""
    p = doc.add_paragraph()
    if level == 1:
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)
        p.paragraph_format.space_before = Pt(24)
        p.paragraph_format.space_after = Pt(12)
    elif level == 2:
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x2D, 0x2D, 0x2D)
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)

def add_screen(label):
    """标注屏幕展示内容"""
    p = doc.add_paragraph()
    run = p.add_run(f'📺 屏幕展示：{label}')
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x00, 0x80, 0x80)

def add_speak(text):
    """口头描述段落"""
    p = doc.add_paragraph()
    run = p.add_run(f'🗣️ {text}')
    run.font.size = Pt(11)

def add_separator():
    doc.add_paragraph('─' * 50)

# ===== 正文 =====

add_section_title('第一部分：打开仓库，讲清楚"我们在做什么"（约2分钟）')

add_screen('打开 GitHub 仓库主页 https://github.com/fly69783/ai-vision-assistant')

add_speak(
    '我们的项目叫做"AI 视觉辅助盲人环境理解系统"。简单来说，就是让视障用户用手机摄像头'
    '拍一下周围环境，系统识别出关键物体和文字，然后用简短的语音告诉他——'
    '比如"前面是走廊，右侧有一扇门，地上没有什么障碍物"。'
)
add_speak(
    '这不是导航系统，是环境理解辅助工具。我们不做测距，不做人脸识别。'
)

add_separator()

add_section_title('第二部分：打开 README.md，讲核心流程（约3分钟）')

add_screen('滚动 README.md，停在"一句话方案"部分')

add_speak(
    '这是我们的核心流程：摄像头拍一张照片 → 先检查画面质量（太模糊或太暗就提示用户重拍）'
    '→ 然后三个 AI 模块同时工作——目标检测找物体、OCR 读文字、视觉语言模型理解场景'
    '→ 把三路结果融合排序 → 最后用简短的语音播报出来。'
)

add_speak(
    '比赛的优秀作品一般都会讲"规划-执行-反馈-优化"的闭环，我们也画了这个闭环：'
    '用户提出任务 → 分析 → 播报 → 如果失败就分类记录 → 修改模型或规则 → 重新测试验证。'
)

add_screen('停在"设计亮点"表格')

add_speak(
    '我们的设计亮点有四个：\n'
    '第一，多模型融合而不是只依赖一个 AI——单一模型容易漏东西或者胡说八道，'
    '我们让检测、OCR、视觉模型各自提供证据，然后交叉验证。\n'
    '第二，面向听觉优化——普通 AI 看图说话可能说一大段，但盲人用户听不了那么长。'
    '我们最多 3 句话，重要的先说。\n'
    '第三，安全降级——画面模糊、光线暗、网络断了，系统会明确提示"不确定"，而不是瞎猜。\n'
    '第四，时序去重——看视频的时候，同一个东西不会反复播报。'
)

add_separator()

add_section_title('第三部分：打开 PROJECT_PLAN.md，讲计划安排（约2分钟）')

add_screen('滚动到"8周完整计划表"（第7节）')

add_speak(
    '我们的 8 周安排是这样的：\n'
    '第 1-2 周是打基础——跑通目标检测、OCR、视觉问答、语音合成的独立示例，'
    '同时收集 40 个测试场景。\n'
    '第 3-4 周是把流程串起来——做个后端接口，上传图片能返回分析结果，然后语音播报。'
    '第 4 周末要能端到端连续演示 3 次。\n'
    '第 5-6 周是完善——加跟踪去重、异常处理、交互界面，做两轮测试。\n'
    '第 7-8 周是收尾——不再加新功能，专注优化稳定性、准备 PPT 和视频、模拟答辩。\n'
    '这是我们现在在第 1 周的状态——环境已经搭好，仓库也建好了，接下来要开始写第一个代码里程碑：OpenCV 读取图片。'
)

add_screen('滚动到分工表（第6节）')

add_speak(
    '分工上，队友主要负责技术开发，我这边负责测试数据整理、部分辅助模块，以及文档和材料。'
    '我们每周末会合在一起演示这周可运行的内容，录 1 分钟视频。'
)

add_separator()

add_section_title('第四部分：打开"项目实施与学习大纲.md"，补充关键信息（约2分钟）')

add_screen('停在"典型用户流程"（第三节）')

add_speak(
    '这是四个典型使用场景：\n'
    '环境概述——用户按一下按钮，系统告诉他前面是什么地方、有什么东西。\n'
    '寻找物品——用户说"门在哪里"，系统只关注门，告诉他大致方向。\n'
    '读取文字——用户说"读一下前面的字"，系统识别文字并按阅读顺序播报。\n'
    '连续观察——摄像头持续低频率取帧，只有发生重要变化才再次提醒。\n'
    '每个场景我们都设计了具体的输入输出示例。'
)

add_screen('停在"安全边界"相关内容（第一节第3点或第十四节）')

add_speak(
    '最后要强调安全边界：我们不替代盲杖和导盲犬，不做精确测距和道路安全决策，'
    '不做人脸识别，数据不保存敏感原图。这是我们诚实的定位。'
)

add_separator()

add_section_title('第五部分：总结收尾（约1分钟）')

add_speak(
    '总结一下：\n'
    '1. 我们的作品是一个辅助工具，帮视障用户理解周围环境，不是导航系统。\n'
    '2. 技术路线上，我们用多模型融合而不是单一模型，优先安全和可信。\n'
    '3. 时间线上，第 4 周出 MVP，第 8 周提交完整材料。\n'
    '4. 我们有一些设计亮点（融合排序、面向听觉、安全降级），但这些目前还是设计方案，'
    '还没实现——整个项目现在处于第 1 周的起步阶段。\n'
    '有什么问题可以随时问我。'
)

add_separator()

add_section_title('附：应急应对 —— 如果被问到不懂的技术问题')

add_speak(
    '如果被问到技术细节不太懂的：\n'
    '"这个部分目前主要由队友在负责设计方案，我还在学习阶段。'
    '我的理解是 [如实说]。如果理解有误，我们会后再核对一下。"\n'
    '这样做反而加分——评委看得出你有没有在伪造。'
)

add_separator()

add_section_title('附：时间控制总览')

# 添加时间表
table = doc.add_table(rows=8, cols=3, style='Light Grid Accent 1')
headers = ['部分', '内容', '预计时间']
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = h
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

data = [
    ['第一部分', '打开仓库，讲清楚项目是什么', '2 分钟'],
    ['第二部分', '打开 README.md，讲核心流程', '3 分钟'],
    ['第三部分', '打开 PROJECT_PLAN.md，讲计划安排', '2 分钟'],
    ['第四部分', '打开学习大纲，补充场景和安全边界', '2 分钟'],
    ['第五部分', '总结收尾', '1 分钟'],
    ['应急', '预留问答缓冲', '2-3 分钟'],
    ['合计', '', '约 12-13 分钟'],
]
for i, row_data in enumerate(data):
    for j, text in enumerate(row_data):
        table.rows[i + 1].cells[j].text = text

# ===== 保存 =====
output_dir = r'E:\竞赛\2027计设'
os.makedirs(output_dir, exist_ok=True)
output_path = os.path.join(output_dir, 'AI视觉辅助系统-项目介绍演讲稿.docx')
doc.save(output_path)
print(f'Done. Document saved to: {output_path}')
